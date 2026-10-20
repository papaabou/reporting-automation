"""Génération de la présentation PowerPoint (16:9, 6 diapositives)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .extract import ReportData
from .theme import BLEU_FONCE, BLEU_MOYEN, GRIS_CLAIR, ORANGE, ROUGE_KO, VERT_OK


def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


def eur(v: float) -> str:
    return f"{v:,.0f} €".replace(",", " ")


def pct(v: float, signed: bool = False) -> str:
    return (f"{v * 100:+.1f} %" if signed else f"{v * 100:.1f} %").replace(".", ",")


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _title(slide, text: str) -> None:
    tf = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.85)).text_frame
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(30), True, "Arial"
    run.font.color.rgb = _rgb(BLEU_FONCE)


def _table(slide, header: list[str], rows: list[list[str]], x: float, y: float, width: float, row_h: float = 0.4) -> None:
    data = [header, *rows]
    table = slide.shapes.add_table(len(data), len(header), Inches(x), Inches(y), Inches(width), Inches(row_h * len(data))).table
    for i, line in enumerate(data):
        for j, value in enumerate(line):
            cell = table.cell(i, j)
            cell.text = str(value)
            p = cell.text_frame.paragraphs[0]
            p.font.size, p.font.name = Pt(11 if len(data) <= 9 else 10), "Arial"
            cell.fill.solid()
            if i == 0:
                p.font.bold = True
                p.font.color.rgb = _rgb("FFFFFF")
                cell.fill.fore_color.rgb = _rgb(BLEU_FONCE)
            else:
                p.font.color.rgb = _rgb("1A1A1A")
                cell.fill.fore_color.rgb = _rgb(GRIS_CLAIR if i % 2 == 0 else "FFFFFF")
            if j > 0:
                p.alignment = PP_ALIGN.RIGHT


def write_pptx(d: ReportData, path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    k = d.kpis

    # 1. Couverture
    s = _blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(BLEU_FONCE)
    bg.line.fill.background()
    tf = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.3), Inches(2.2)).text_frame
    run = tf.paragraphs[0].add_run()
    run.text = d.titre
    run.font.size, run.font.bold, run.font.name = Pt(44), True, "Arial"
    run.font.color.rgb = _rgb("FFFFFF")
    run = tf.add_paragraph().add_run()
    run.text = f"Rapport mensuel · {d.periode}"
    run.font.size, run.font.name = Pt(24), "Arial"
    run.font.color.rgb = _rgb("F0A868")

    # 2. Indicateurs clés
    s = _blank(prs)
    _title(s, f"Indicateurs clés — {d.periode}")
    cards = [
        (eur(k["ventes_mois"]), "Ventes du mois", BLEU_FONCE),
        (pct(k["realisation_mois"]), "Réalisation du mois", VERT_OK if k["realisation_mois"] >= 1 else ORANGE),
        (eur(k["ventes_ytd"]), "Ventes cumulées", BLEU_FONCE),
        (pct(k["realisation_ytd"]), "Réalisation cumulée", VERT_OK if k["realisation_ytd"] >= 1 else ORANGE),
        (pct(k["croissance_ytd"], signed=True), "Croissance vs N-1", VERT_OK if k["croissance_ytd"] >= 0 else ROUGE_KO),
    ]
    for i, (value, label, color) in enumerate(cards):
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + i * 2.52), Inches(2.1), Inches(2.3), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb("FFFFFF")
        card.line.color.rgb = _rgb("D5DEE5")
        card.shadow.inherit = False
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.35)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = value
        run.font.size, run.font.bold, run.font.name = Pt(22), True, "Arial"
        run.font.color.rgb = _rgb(color)
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size, run.font.name = Pt(12), "Arial"
        run.font.color.rgb = _rgb("555555")

    # 3. Graphique ventes vs budget
    s = _blank(prs)
    _title(s, "Ventes mensuelles vs budget (€)")
    chart_data = CategoryChartData()
    chart_data.categories = [m["mois"] for m in d.mensuel]
    chart_data.add_series("Budget N", [m["budget"] for m in d.mensuel])
    chart_data.add_series("Ventes N", [m["ventes"] for m in d.mensuel])
    chart_data.add_series("Ventes N-1", [m["ventes_n1"] for m in d.mensuel])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.35), Inches(12.1), Inches(5.7), chart_data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    for series, color in zip(chart.plots[0].series, [BLEU_MOYEN, ORANGE, "9AA8B2"]):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _rgb(color)

    # 4. Par équipe
    s = _blank(prs)
    _title(s, f"Performance par équipe — {d.periode}")
    _table(s, ["Équipe", "Ventes N-1", "Budget N", "Ventes N", "Réalisation", "Croissance"],
           [[e["equipe"], eur(e["ventes_n1"]), eur(e["budget"]), eur(e["ventes"]), pct(e["realisation"]), pct(e["croissance"], True)]
            for e in d.equipes], 0.6, 1.5, 12.1, 0.5)

    # 5. Top produits
    s = _blank(prs)
    _title(s, f"Top {len(d.produits)} produits (ventes N)")
    _table(s, ["Produit", "Budget N", "Ventes N-1", "Ventes N"],
           [[p["produit"], eur(p["budget"]), eur(p["ventes_n1"]), eur(p["ventes"])] for p in d.produits], 0.6, 1.35, 12.1, 0.48)

    # 6. Stock
    s = _blank(prs)
    _title(s, f"Stock & commandes par grossiste — {d.periode}")
    _table(s, ["Grossiste", "Ventes (Un)", "Stock (Un)", "Commandes en attente (Un)"],
           [[g["grossiste"], f"{g['ventes_un']:,.0f}".replace(",", " "), f"{g['stock_un']:,.0f}".replace(",", " "),
             f"{g['commandes_un']:,.0f}".replace(",", " ")] for g in d.stock], 2.3, 1.8, 8.7, 0.55)

    path = Path(path)
    prs.save(path)
    return path
