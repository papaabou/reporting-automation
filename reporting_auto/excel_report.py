"""Génération du classeur Excel : vrais Tableaux Excel, formats € / %, totaux en formules."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .extract import ReportData
from .theme import BLEU_FONCE, GRIS_CLAIR

EUR, PCT, UN = '#,##0 "€"', "0.0%", "#,##0"


def _write_table(ws, header: list[str], rows: list[list], name: str, formats: dict[int, str], total: bool = False) -> None:
    thin = Side(style="thin", color="C9D4DC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    data = [header, *rows]

    for i, line in enumerate(data, start=1):
        for j, value in enumerate(line, start=1):
            c = ws.cell(row=i, column=j, value=value)
            c.font = Font(name="Arial", size=10)
            c.border = border
            if i == 1:
                c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=BLEU_FONCE)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif j in formats:
                c.number_format = formats[j]

    n_rows, n_cols = len(data), len(header)
    if total:
        r = n_rows + 1
        ws.cell(row=r, column=1, value="TOTAL").font = Font(name="Arial", size=10, bold=True)
        for j in range(2, n_cols + 1):
            if formats.get(j) == PCT:  # on ne somme pas des pourcentages
                continue
            col = get_column_letter(j)
            c = ws.cell(row=r, column=j, value=f"=SUM({col}2:{col}{n_rows})")
            c.font = Font(name="Arial", size=10, bold=True)
            c.fill = PatternFill("solid", fgColor=GRIS_CLAIR)
            c.number_format = formats.get(j, "General")

    table = Table(displayName=name, ref=f"A1:{get_column_letter(n_cols)}{n_rows}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    for j in range(1, n_cols + 1):
        width = max(len(f"{line[j - 1]:,.0f}" if isinstance(line[j - 1], float) else str(line[j - 1])) for line in data)
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 3, 12), 42)


def write_excel(d: ReportData, path: str | Path) -> Path:
    wb = Workbook()

    ws = wb.active
    ws.title = "Ventes mensuelles"
    _write_table(ws, ["Mois", "Ventes N-2 (€)", "Ventes N-1 (€)", "Budget N (€)", "Ventes N (€)", "Réalisation (%)"],
                 [[m["mois"], m["ventes_n2"], m["ventes_n1"], m["budget"], m["ventes"], m["realisation"]] for m in d.mensuel],
                 "TabMensuel", {2: EUR, 3: EUR, 4: EUR, 5: EUR, 6: PCT}, total=True)

    ws = wb.create_sheet("Cumul")
    _write_table(ws, ["Mois", "Cumul N-2 (€)", "Cumul N-1 (€)", "Budget cumulé N (€)", "Cumul N (€)"],
                 [[c["mois"], c["ventes_n2"], c["ventes_n1"], c["budget"], c["ventes"]] for c in d.cumul],
                 "TabCumul", {2: EUR, 3: EUR, 4: EUR, 5: EUR})

    ws = wb.create_sheet("Trimestres")
    _write_table(ws, ["Trimestre", "Ventes N-2 (€)", "Ventes N-1 (€)", "Budget N (€)", "Ventes N (€)"],
                 [[t["trimestre"], t["ventes_n2"], t["ventes_n1"], t["budget"], t["ventes"]] for t in d.trimestres],
                 "TabTrimestres", {2: EUR, 3: EUR, 4: EUR, 5: EUR}, total=True)

    ws = wb.create_sheet("Par équipe")
    _write_table(ws, ["Équipe", "Ventes N-1 (€)", "Budget N (€)", "Ventes N (€)", "Réalisation (%)", "Croissance (%)"],
                 [[e["equipe"], e["ventes_n1"], e["budget"], e["ventes"], e["realisation"], e["croissance"]] for e in d.equipes],
                 "TabEquipes", {2: EUR, 3: EUR, 4: EUR, 5: PCT, 6: PCT}, total=True)

    ws = wb.create_sheet("Top produits")
    _write_table(ws, ["Produit", "Budget N (€)", "Ventes N-1 (€)", "Ventes N (€)"],
                 [[p["produit"], p["budget"], p["ventes_n1"], p["ventes"]] for p in d.produits],
                 "TabProduits", {2: EUR, 3: EUR, 4: EUR})

    ws = wb.create_sheet("Stock & commandes")
    _write_table(ws, ["Grossiste", "Ventes (Un)", "Stock (Un)", "Commandes en attente (Un)"],
                 [[s["grossiste"], s["ventes_un"], s["stock_un"], s["commandes_un"]] for s in d.stock],
                 "TabStock", {2: UN, 3: UN, 4: UN}, total=True)

    path = Path(path)
    wb.save(path)
    return path
