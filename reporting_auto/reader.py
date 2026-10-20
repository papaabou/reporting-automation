"""Lecture uniforme des classeurs .xlsb (binaire Excel) et .xlsx."""
from __future__ import annotations

import unicodedata
from pathlib import Path

Rows = list[list]


def read_workbook(path: str | Path) -> dict[str, Rows]:
    """Retourne {nom de feuille: lignes}, quel que soit le format du classeur."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsb":
        from pyxlsb import open_workbook

        sheets = {}
        with open_workbook(str(path)) as wb:
            for name in wb.sheets:
                with wb.get_sheet(name) as sh:
                    sheets[name] = [[c.v for c in row] for row in sh.rows()]
        return sheets
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        wb.close()
        return sheets
    raise ValueError(f"Format non pris en charge : {suffix} (attendu .xlsb ou .xlsx)")


def normalize(text) -> str:
    """Minuscules, sans accents ni espaces superflus : pour comparer des intitulés."""
    if text is None:
        return ""
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return " ".join(text.lower().split())


def find_sheet(sheets: dict[str, Rows], keyword: str) -> str:
    key = normalize(keyword)
    for name in sheets:
        if key in normalize(name):
            return name
    raise KeyError(f"Aucune feuille ne contient « {keyword} » (feuilles : {', '.join(sheets)})")


def extract_table(rows: Rows, columns: dict[str, list[str]]) -> list[dict]:
    """
    Repère la ligne d'en-tête qui contient toutes les colonnes demandées
    (en acceptant plusieurs intitulés possibles), puis lit les lignes suivantes
    jusqu'à la première ligne vide.
    """
    aliases = {field: {normalize(a) for a in names} for field, names in columns.items()}
    for h, header in enumerate(rows):
        cells = [normalize(c) for c in header]
        positions = {}
        for field, names in aliases.items():
            idx = next((i for i, c in enumerate(cells) if c in names), None)
            if idx is None:
                break
            positions[field] = idx
        else:
            records = []
            for row in rows[h + 1:]:
                if not row or all(v in (None, "") for v in row):
                    break
                records.append({f: (row[i] if i < len(row) else None) for f, i in positions.items()})
            return records
    expected = ", ".join(names[0] for names in columns.values())
    raise KeyError(f"En-tête introuvable (colonnes attendues : {expected})")
