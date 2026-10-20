"""
Génère un classeur de démonstration avec des données 100 % fictives,
volontairement « brut » (feuille menu, lignes d'en-tête décalées, lignes de total)
pour reproduire les difficultés d'un vrai export.
"""
from __future__ import annotations

import random
from pathlib import Path

from openpyxl import Workbook

from .config import MOIS

EQUIPES = ["Nord", "Sud", "Est", "Ouest", "Centre", "Grands comptes"]
PRODUITS = [
    "Gamme Alpha 50", "Gamme Alpha 100", "Beta Plus", "Beta Kids", "Gamma Forte", "Delta Spray",
    "Epsilon Gel", "Zeta 20", "Eta Duo", "Theta Max", "Iota Sirop", "Kappa Daily", "Lambda Night", "Mu Care",
]
GROSSISTES = ["Distrib A", "Distrib B", "Distrib C", "Distrib D"]


def generate_sample(path: str | Path, mois_clos: int = 3, seed: int = 7) -> Path:
    """Crée un classeur fictif dont les `mois_clos` premiers mois de l'exercice sont réalisés."""
    rng = random.Random(seed)
    wb = Workbook()

    ws = wb.active
    ws.title = "Menu"
    ws.append(["Tableau de bord des ventes — DONNÉES FICTIVES"])
    ws.append([])
    ws.append(["Période :", f"{MOIS[mois_clos - 1]} 26"])

    ws = wb.create_sheet("Ventes mensuelles")
    ws.append(["Export du système de ventes"])
    ws.append([])
    ws.append(["Mois", "Ventes N-2", "Ventes N-1", "Budget N", "Ventes N"])
    base = 180_000
    for i, m in enumerate(MOIS):
        season = 1 + 0.15 * ((i % 6) - 2.5) / 2.5
        n2 = base * season * rng.uniform(0.9, 1.05)
        n1 = n2 * rng.uniform(1.02, 1.12)
        budget = n1 * 1.08
        ventes = budget * rng.uniform(0.9, 1.12) if i < mois_clos else None
        ws.append([m, round(n2), round(n1), round(budget), round(ventes) if ventes else None])

    ws = wb.create_sheet("Ventes par equipe")
    ws.append(["Équipe", "Ventes N-1", "Budget N", "Ventes N"])
    for e in EQUIPES:
        n1 = rng.uniform(60_000, 160_000)
        budget = n1 * 1.08
        ws.append([e, round(n1), round(budget), round(budget * rng.uniform(0.85, 1.15))])
    ws.append(["Total", None, None, None])

    ws = wb.create_sheet("Report by Produit")
    ws.append(["Rapport par produit"])
    ws.append(["Code", "Produit", "Budget N", "Ventes N-1", "Ventes N"])
    for k, p in enumerate(PRODUITS):
        n1 = rng.uniform(10_000, 90_000)
        budget = n1 * 1.08
        ws.append([f"SKU-{100 + k}", p, round(budget), round(n1), round(budget * rng.uniform(0.8, 1.2))])

    ws = wb.create_sheet("Stock")
    ws.append(["Grossiste", "Produit", "Ventes (Un)", "Stock (Un)", "Commandes en attente (Un)"])
    for g in GROSSISTES:
        for p in PRODUITS:
            ws.append([g, p, rng.randint(50, 900), rng.randint(0, 1500), rng.randint(0, 300)])

    path = Path(path)
    wb.save(path)
    return path
