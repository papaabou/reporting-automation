import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from reporting_auto import extract, generate_sample, run
from reporting_auto.reader import extract_table, normalize


@pytest.fixture
def sample(tmp_path):
    return generate_sample(tmp_path / "demo.xlsx", mois_clos=3)


def test_extract_reads_all_tables(sample):
    d = extract(sample)
    assert d.periode == "Juin 26"
    assert [m["mois"] for m in d.mensuel][:3] == ["Avr", "Mai", "Juin"]
    assert len(d.mois_clos) == 3
    assert len(d.equipes) == 6            # la ligne « Total » est ignorée
    assert len(d.produits) == 10          # top 10
    assert d.produits[0]["ventes"] >= d.produits[-1]["ventes"]
    assert len(d.stock) == 4              # agrégé par grossiste


def test_kpis_are_consistent(sample):
    d = extract(sample)
    k = d.kpis
    assert k["ventes_ytd"] == pytest.approx(sum(m["ventes"] for m in d.mensuel))
    assert k["ventes_mois"] == d.mois_clos[-1]["ventes"]
    assert sum(t["budget"] for t in d.trimestres) == pytest.approx(sum(m["budget"] for m in d.mensuel))


def test_header_detection_accepts_aliases_and_accents():
    rows = [["titre"], [], ["Team", "Sales N-1", "Budget", "Sales N"], ["Nord", 1, 2, 3], []]
    cols = {"equipe": ["équipe", "team"], "ventes_n1": ["ventes n-1", "sales n-1"],
            "budget": ["budget"], "ventes": ["ventes n", "sales n"]}
    assert extract_table(rows, cols) == [{"equipe": "Nord", "ventes_n1": 1, "budget": 2, "ventes": 3}]
    assert normalize("  Équipe  Nord ") == "equipe nord"


def test_missing_sheet_gives_clear_error(tmp_path):
    wb = Workbook()
    wb.active.title = "Autre"
    wb.save(tmp_path / "vide.xlsx")
    with pytest.raises(KeyError, match="Aucune feuille"):
        extract(tmp_path / "vide.xlsx")


def test_full_pipeline_outputs(sample, tmp_path):
    xlsx, pptx = run(sample, tmp_path / "out", log=lambda *_: None)
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Ventes mensuelles", "Cumul", "Trimestres", "Par équipe", "Top produits", "Stock & commandes"]
    assert str(wb["Ventes mensuelles"]["B14"].value).startswith("=SUM(")   # totaux en formules
    assert len(Presentation(pptx).slides) == 6
